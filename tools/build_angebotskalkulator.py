#!/usr/bin/env python3
"""Erzeugt den Angebotskalkulator (Handwerk & Logistik) als .xlsx.

Aufruf:  python3 tools/build_angebotskalkulator.py [zielpfad.xlsx]

Alle Rechenwege stehen als echte Excel-Formeln in der Datei, damit sich das
Angebot bei geänderten Eingaben neu berechnet. Die Formelkette ist in
docs/rechenmodell.md beschrieben und ist zugleich die Referenz für die
Angebots-Engine der App.
"""

import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --- Gestaltung -----------------------------------------------------------
FONT = "Arial"

BLUE = "0000FF"      # harte Eingabewerte
BLACK = "000000"     # Formeln
GREEN = "008000"     # Verweise auf ein anderes Blatt
YELLOW = "FFFF00"    # Schlüsselannahmen / auszufüllen

H1 = Font(name=FONT, size=14, bold=True)
H2 = Font(name=FONT, size=11, bold=True, color="FFFFFF")
LABEL = Font(name=FONT, size=10)
LABEL_B = Font(name=FONT, size=10, bold=True)
INPUT_F = Font(name=FONT, size=10, color=BLUE)
CALC_F = Font(name=FONT, size=10, color=BLACK)
LINK_F = Font(name=FONT, size=10, color=GREEN)
NOTE_F = Font(name=FONT, size=9, italic=True, color="606060")

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
BAND_FILL = PatternFill("solid", fgColor="DCE6F1")
KEY_FILL = PatternFill("solid", fgColor=YELLOW)
TOTAL_FILL = PatternFill("solid", fgColor="D9D9D9")

THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_LINE = Border(top=Side(style="medium", color="1F3864"))

EUR = '#,##0.00" €"'
EUR0 = '#,##0" €"'
NUM2 = '#,##0.00'
NUM3 = '#,##0.000'
PCT = '0.0%'

# --- Stammdaten: Kalkulationsgrößen ------------------------------------
# Zeile in "Stammdaten" -> (Bezeichnung, Wert, Einheit, Hinweis)
FIRMA = [
    ("Firmenname", "Muster Malerbetrieb GmbH"),
    ("Straße, Nr.", "Handwerkerweg 12"),
    ("PLZ / Ort", "12345 Musterstadt"),
    ("Telefon", "0123 456789"),
    ("E-Mail", "info@muster-malerbetrieb.de"),
    ("USt-IdNr.", "DE123456789"),
    ("IBAN", "DE00 0000 0000 0000 0000 00"),
]

KALK = [
    ("Mittellohn (Kosten je produktiver Stunde)", 42.00, EUR,
     "Bruttolohn + Lohnnebenkosten / produktive Stunden. Richtwert - durch eigenen Wert ersetzen."),
    ("Gemeinkostenzuschlag auf Lohn", 0.25, PCT,
     "Büro, Fahrzeuge, Werkzeug, Versicherung, unproduktive Zeiten."),
    ("Gemeinkostenzuschlag auf Material", 0.10, PCT,
     "Beschaffung, Lagerung, Verschnitt."),
    ("Wagnis & Gewinn", 0.12, PCT,
     "Aufschlag auf die Selbstkosten. Handwerk üblich 8-15 %."),
    ("Rabatt (auf Angebotssumme)", 0.00, PCT,
     "Optionaler Nachlass. 0 = kein Rabatt."),
    ("Skonto (einkalkuliert)", 0.02, PCT,
     "Wird auf den Preis aufgeschlagen, damit ein gezogenes Skonto die Marge nicht auffrisst."),
    ("Fahrtkosten je km", 0.70, EUR,
     "Kilometersatz Nutzfahrzeug inkl. Verschleiß."),
    ("Umsatzsteuersatz", 0.19, PCT,
     "19 % Regelsatz. Bei Kleinunternehmer nach § 19 UStG auf 0 setzen."),
    ("Zahlungsziel (Tage)", 14, "0",
     "Erscheint im Angebotstext."),
    ("Angebot gültig (Tage)", 30, "0",
     "Bindefrist ab Angebotsdatum."),
]

# --- Leistungskatalog -----------------------------------------------------
# Nr, Gewerk, Leistung, Einheit, Zeit/Einheit (Std), Material/Einheit (EUR), Hinweis
KATALOG = [
    ("M01", "Maler", "Böden und Mobiliar abdecken, Kanten abkleben", "m²", 0.05, 0.35, "Bezug: Bodenfläche"),
    ("M02", "Maler", "Untergrund reinigen und vorbereiten", "m²", 0.04, 0.10, ""),
    ("M03", "Maler", "Risse und Löcher spachteln, schleifen", "m²", 0.06, 0.45, "Nur anteilige Fläche ansetzen"),
    ("M04", "Maler", "Tiefengrund / Haftgrund auftragen", "m²", 0.03, 0.55, ""),
    ("M05", "Maler", "Wände streichen, 2 Anstriche Dispersion", "m²", 0.12, 1.25, "Farbe ca. 0,35 l/m²"),
    ("M06", "Maler", "Decken streichen, 2 Anstriche Dispersion", "m²", 0.15, 1.25, "Über Kopf, höherer Zeitansatz"),
    ("M07", "Maler", "Fenster- und Heizkörpernischen streichen", "m²", 0.25, 1.60, "Kleinteilig"),
    ("M08", "Maler", "Türblatt lackieren, beidseitig", "Stk", 1.20, 12.00, ""),
    ("M09", "Maler", "Türzarge lackieren", "Stk", 0.80, 7.50, ""),
    ("M10", "Maler", "Endreinigung und Abfallentsorgung", "m²", 0.04, 0.25, "Bezug: Bodenfläche"),
    ("M11", "Maler", "Raufasertapete entfernen", "m²", 0.18, 0.30, ""),
    ("M12", "Maler", "An- und Abfahrt, Rüst- und Räumzeit", "pausch.", 1.50, 0.00, "Je Einsatz"),
    ("T01", "Trockenbau", "Gipskartonwand beplanken, einlagig", "m²", 0.35, 14.50, ""),
    ("T02", "Trockenbau", "Fugen spachteln Q2", "m²", 0.12, 1.10, ""),
    ("B01", "Boden", "Laminat verlegen inkl. Trittschall", "m²", 0.25, 18.00, ""),
    ("B02", "Boden", "Sockelleisten montieren", "m", 0.10, 4.20, ""),
    ("E01", "Elektro", "Steckdose / Schalter tauschen", "Stk", 0.35, 6.50, ""),
    ("E02", "Elektro", "Leuchte montieren und anschließen", "Stk", 0.50, 0.00, "Leuchte bauseits"),
    ("L01", "Logistik", "Be- und Entladen (Personalstunde)", "Std", 1.00, 0.00, ""),
    ("L02", "Logistik", "Transportfahrt Nutzfahrzeug", "km", 0.02, 0.45, "Zeit + Betriebskosten je km"),
    ("L03", "Logistik", "Umzugskarton packen inkl. Material", "Stk", 0.12, 2.20, ""),
    ("L04", "Logistik", "Möbeldemontage / -montage", "Std", 1.00, 0.00, ""),
    ("L05", "Logistik", "Einlagerung je m³ und Monat", "m³", 0.05, 8.00, ""),
    ("L06", "Logistik", "Halteverbotszone beantragen und stellen", "pausch.", 1.00, 95.00, ""),
]

# Vorbelegte Beispielkalkulation: 45 m² Wohnung komplett streichen.
# (Katalog-Nr, Mengenformel oder Zahl)
BEISPIEL = [
    ("M12", 1),
    ("M01", "=$B$18"),
    ("M02", "=$B$16+$B$17"),
    ("M03", "=ROUND(($B$16+$B$17)*$B$20,1)"),
    ("M04", "=$B$16+$B$17"),
    ("M05", "=$B$16"),
    ("M06", "=$B$17"),
    ("M09", "=$B$19"),
    ("M08", "=$B$19"),
    ("M10", "=$B$18"),
]

POS_FIRST = 25          # erste Positionszeile in "Kalkulation"
POS_ROWS = 20           # Anzahl Positionszeilen
POS_LAST = POS_FIRST + POS_ROWS - 1
KAT_FIRST, KAT_LAST = 4, 4 + len(KATALOG) - 1


def set_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def title(ws, text, span):
    ws["A1"] = text
    ws["A1"].font = H1
    ws.merge_cells(f"A1:{span}1")
    ws.row_dimensions[1].height = 22


def section(ws, row, text, span):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = H2
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = HEAD_FILL
    return cell


# =========================================================================
def build_anleitung(wb):
    ws = wb.create_sheet("Anleitung")
    set_widths(ws, {"A": 22, "B": 96})
    title(ws, "Angebotskalkulator - Handwerk & Logistik", "B")

    rows = [
        ("", ""),
        ("Was die Datei tut",
         "Aus Aufmaß und Leistungen berechnet sie Stunden, Material, Selbstkosten und daraus "
         "einen belastbaren Angebotspreis - inklusive druckfertigem Angebot."),
        ("", ""),
        ("Farblegende", ""),
        ("blaue Schrift", "Eingabefeld - hier tragen Sie Ihre Werte ein."),
        ("schwarze Schrift", "Formel - bitte nicht überschreiben."),
        ("grüne Schrift", "Verweis auf ein anderes Tabellenblatt."),
        ("gelbe Füllung", "Schlüsselannahme - vor dem ersten Angebot einmal prüfen und anpassen."),
        ("", ""),
        ("In dieser Reihenfolge", ""),
        ("1. Stammdaten",
         "Firmendaten eintragen und die Kalkulationsgrößen auf den eigenen Betrieb einstellen. "
         "Wichtigster Wert: der Mittellohn."),
        ("2. Leistungskatalog",
         "Eigene Leistungen ergänzen: Zeit je Einheit und Materialkosten je Einheit. "
         "Die Katalog-Nr. ist der Schlüssel, über den die Kalkulation zugreift."),
        ("3. Kalkulation",
         "Kunde und Aufmaß eintragen, dann je Position eine Katalog-Nr. wählen und die Menge "
         "setzen. Mengen dürfen auf den Aufmaß-Block verweisen, z. B. =B16 für die Wandfläche."),
        ("4. Angebot",
         "Wird automatisch gefüllt und ist auf A4 druckfertig eingerichtet - drucken oder als "
         "PDF exportieren."),
        ("", ""),
        ("Beispiel enthalten",
         "Vorbelegt ist 'Wohnung 45 m² komplett streichen' inkl. Vorarbeiten und Türen. "
         "Einfach die Zahlen im Aufmaß-Block überschreiben - alles rechnet sich neu."),
        ("", ""),
        ("Rechenweg je Position",
         "Stunden = Menge x Zeit/Einheit  |  Lohnkosten = Stunden x Mittellohn  |  "
         "Material = Menge x Material/Einheit"),
        ("",
         "Selbstkosten = Lohnkosten x (1 + GK Lohn) + Material x (1 + GK Material)"),
        ("",
         "Einheitspreis = Selbstkosten x (1 + Wagnis & Gewinn) x (1 + Skonto) / Menge"),
        ("", ""),
        ("Rechenweg Summe",
         "Positionen + An-/Abfahrt - Rabatt = Nettobetrag; zzgl. USt = Gesamtbetrag brutto."),
        ("Warum Skonto im Preis",
         "Der Skontosatz steckt im Einheitspreis statt in einer eigenen Summenzeile. So geht das "
         "gedruckte Angebot für den Kunden sauber auf, und ein gezogenes Skonto kostet keine Marge."),
        ("", ""),
        ("Hinweis zu den Vorgabewerten",
         "Zeit- und Materialansätze im Katalog sind Branchen-Richtwerte zur Orientierung, keine "
         "verbindliche Kalkulation. Maßgeblich sind die Werte des eigenen Betriebs."),
    ]
    r = 3
    for label, text in rows:
        ws.cell(row=r, column=1, value=label).font = LABEL_B if label else LABEL
        c = ws.cell(row=r, column=2, value=text)
        c.font = LABEL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if text and len(text) > 90:
            ws.row_dimensions[r].height = 30
        r += 1

    for row, color in (("7", "0000FF"), ("8", "000000"), ("9", "008000")):
        ws[f"A{row}"].font = Font(name=FONT, size=10, bold=True, color=color)
    ws["A10"].fill = KEY_FILL
    return ws


# =========================================================================
def build_stammdaten(wb):
    ws = wb.create_sheet("Stammdaten")
    set_widths(ws, {"A": 46, "B": 22, "C": 62})
    title(ws, "Stammdaten und Kalkulationsgrößen", "C")

    section(ws, 3, "Firmendaten (erscheinen im Angebotskopf)", 3)
    r = 4
    for label, value in FIRMA:
        ws.cell(row=r, column=1, value=label).font = LABEL
        c = ws.cell(row=r, column=2, value=value)
        c.font = INPUT_F
        c.fill = KEY_FILL
        c.border = BOX
        r += 1

    r += 2
    section(ws, r, "Kalkulationsgrößen", 3)
    ws.cell(row=r, column=2).font = H2
    ws.cell(row=r, column=3).font = H2
    r += 1
    kalk_start = r
    for label, value, fmt, note in KALK:
        ws.cell(row=r, column=1, value=label).font = LABEL
        c = ws.cell(row=r, column=2, value=value)
        c.font = INPUT_F
        c.fill = KEY_FILL
        c.border = BOX
        c.number_format = fmt
        n = ws.cell(row=r, column=3, value=note)
        n.font = NOTE_F
        n.alignment = Alignment(wrap_text=True, vertical="center")
        r += 1

    note = ws.cell(row=r + 1, column=1,
                   value="Alle Vorgabewerte sind Branchen-Richtwerte und vom Nutzer zu prüfen. "
                         "Quelle der Zahlen: Vorbelegung dieser Vorlage, keine externe Preisliste.")
    note.font = NOTE_F
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=3)

    return ws, kalk_start


def build_katalog(wb):
    ws = wb.create_sheet("Leistungskatalog")
    set_widths(ws, {"A": 10, "B": 14, "C": 52, "D": 10, "E": 18, "F": 20, "G": 34})
    title(ws, "Leistungskatalog", "G")
    ws["A2"] = ("Katalog-Nr. ist der Schlüssel für das Blatt 'Kalkulation'. "
                "Eigene Leistungen einfach in freien Zeilen ergänzen.")
    ws["A2"].font = NOTE_F
    ws.merge_cells("A2:G2")

    headers = ["Nr.", "Gewerk", "Leistung", "Einheit", "Zeit/Einheit (Std)",
               "Material/Einheit (€)", "Hinweis"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = H2
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
    ws.row_dimensions[3].height = 30

    for i, (nr, gewerk, leistung, einheit, zeit, mat, hinweis) in enumerate(KATALOG):
        r = KAT_FIRST + i
        values = [nr, gewerk, leistung, einheit, zeit, mat, hinweis]
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = INPUT_F if j in (5, 6) else LABEL
            c.border = BOX
            if j == 5:
                c.number_format = NUM3
                c.fill = KEY_FILL
            elif j == 6:
                c.number_format = EUR
                c.fill = KEY_FILL
            elif j == 7:
                c.font = NOTE_F
        if i % 2:
            for j in range(1, 5):
                ws.cell(row=r, column=j).fill = BAND_FILL

    r = KAT_LAST + 2
    ws.cell(row=r, column=1,
            value="Zeit- und Materialansätze sind Richtwerte dieser Vorlage zur Orientierung "
                  "(keine externe Preisliste) und durch die eigenen Nachkalkulationswerte zu ersetzen.").font = NOTE_F
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.freeze_panes = "A4"
    return ws


# =========================================================================
def build_kalkulation(wb):
    ws = wb.create_sheet("Kalkulation")
    set_widths(ws, {"A": 6, "B": 13, "C": 46, "D": 9, "E": 11, "F": 13, "G": 10,
                    "H": 13, "I": 14, "J": 13, "K": 14, "L": 13, "M": 14})
    title(ws, "Angebotskalkulation", "M")

    # --- Kunde -----------------------------------------------------------
    kunde = [
        ("Kunde / Auftraggeber", "Familie Beispiel"),
        ("Straße, Nr.", "Beispielstrasse 7"),
        ("PLZ / Ort", "12345 Musterstadt"),
        ("Bauvorhaben / Objekt", "Wohnung 45 m², komplett streichen"),
        ("Angebotsnummer", "2026-0001"),
        ("Angebotsdatum", "=TODAY()"),
    ]
    section(ws, 2, "Kunde und Vorgang", 13)
    r = 3
    for label, value in kunde:
        ws.cell(row=r, column=1, value=label).font = LABEL
        c = ws.cell(row=r, column=2, value=value)
        c.font = INPUT_F
        c.fill = KEY_FILL
        c.border = BOX
        if label == "Angebotsdatum":
            c.number_format = "DD.MM.YYYY"
            c.font = CALC_F
            c.fill = PatternFill()
        r += 1

    # --- Aufmaß ---------------------------------------------------------
    section(ws, 10, "Aufmaß-Helfer (Eingaben blau, Ergebnisse schwarz)", 13)
    aufmass = [
        ("Wohnfläche (m²)", 45, NUM2, True,
         "Grundfläche der zu bearbeitenden Räume."),
        ("Raumhöhe (m)", 2.50, NUM2, True, "Nur informativ / für eigenes Nachrechnen."),
        ("Faktor Wandfläche je m² Wohnfläche", 2.50, NUM2, True,
         "Faustwert für Wohnungen mit mehreren Räumen: 2,3-2,8. Bei einem einzelnen "
         "großen Raum niedriger ansetzen."),
        ("Wandfläche brutto (m²)", "=ROUND($B$11*$B$13,2)", NUM2, False, ""),
        ("Abzug Fenster und Türen (m²)", 12, NUM2, True, "Summe der nicht zu streichenden Flächen."),
        ("Wandfläche netto (m²)", "=ROUND(MAX(0,$B$14-$B$15),2)", NUM2, False, ""),
        ("Deckenfläche (m²)", "=$B$11", NUM2, False, ""),
        ("Bodenfläche zum Abdecken (m²)", "=$B$11", NUM2, False, ""),
        ("Anzahl Türen (Stk)", 4, "0", True, ""),
        ("Anteil Spachtel-/Ausbesserungsfläche", 0.30, PCT, True,
         "Wie viel Prozent der Fläche muss tatsächlich gespachtelt werden."),
        ("Fahrstrecke gesamt (km, hin und zurück)", 40, "0", True, ""),
    ]
    r = 11
    for label, value, fmt, is_input, note in aufmass:
        ws.cell(row=r, column=1, value=label).font = LABEL
        c = ws.cell(row=r, column=2, value=value)
        c.number_format = fmt
        c.border = BOX
        if is_input:
            c.font = INPUT_F
            c.fill = KEY_FILL
        else:
            c.font = CALC_F
        if note:
            c.comment = Comment(note, "Vorlage")
            n = ws.cell(row=r, column=3, value=note)
            n.font = NOTE_F
        r += 1

    # --- Positionen ------------------------------------------------------
    section(ws, 23, "Positionen", 13)
    headers = [
        "Pos", "Kat.-Nr.", "Leistung", "Einheit", "Menge", "Zeit/Einh. (Std)",
        "Stunden", "Lohnkosten (€)", "Material/Einh. (€)", "Material (€)",
        "Selbstkosten (€)", "EP netto (€)", "GP netto (€)",
    ]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=24, column=i, value=h)
        c.font = H2
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
    ws.row_dimensions[24].height = 32

    kat_nr = f"Leistungskatalog!$A${KAT_FIRST}:$A${KAT_LAST}"
    for i in range(POS_ROWS):
        r = POS_FIRST + i
        katnr, menge = BEISPIEL[i] if i < len(BEISPIEL) else ("", None)

        ws.cell(row=r, column=1, value=i + 1).font = CALC_F
        b = ws.cell(row=r, column=2, value=katnr)
        b.font = INPUT_F
        b.fill = KEY_FILL

        def lookup(col_letter):
            return (f'=IFERROR(INDEX(Leistungskatalog!${col_letter}${KAT_FIRST}:'
                    f'${col_letter}${KAT_LAST},MATCH($B{r},{kat_nr},0)),"")')

        ws.cell(row=r, column=3, value=lookup("C")).font = LINK_F
        d = ws.cell(row=r, column=4, value=lookup("D"))
        d.font = LINK_F
        d.alignment = Alignment(horizontal="center")

        e = ws.cell(row=r, column=5, value=menge if menge is not None else None)
        e.font = INPUT_F
        e.fill = KEY_FILL
        e.number_format = NUM2

        f = ws.cell(row=r, column=6,
                    value=(f'=IFERROR(INDEX(Leistungskatalog!$E${KAT_FIRST}:$E${KAT_LAST},'
                           f'MATCH($B{r},{kat_nr},0)),0)'))
        f.font = LINK_F
        f.number_format = NUM3

        ws.cell(row=r, column=7, value=f'=IF($B{r}="","",ROUND($E{r}*$F{r},2))').number_format = NUM2
        ws.cell(row=r, column=8,
                value=f'=IF($B{r}="","",ROUND($G{r}*Stammdaten!$B$14,2))').number_format = EUR

        i_cell = ws.cell(row=r, column=9,
                         value=(f'=IFERROR(INDEX(Leistungskatalog!$F${KAT_FIRST}:$F${KAT_LAST},'
                                f'MATCH($B{r},{kat_nr},0)),0)'))
        i_cell.font = LINK_F
        i_cell.number_format = EUR

        ws.cell(row=r, column=10, value=f'=IF($B{r}="","",ROUND($E{r}*$I{r},2))').number_format = EUR
        ws.cell(row=r, column=11,
                value=(f'=IF($B{r}="","",ROUND($H{r}*(1+Stammdaten!$B$15)'
                       f'+$J{r}*(1+Stammdaten!$B$16),2))')).number_format = EUR
        # Skonto steckt im Einheitspreis, damit das gedruckte Angebot ohne
        # Zusatzzeile aufgeht: Positionen + Fahrt - Rabatt = Nettobetrag.
        ws.cell(row=r, column=12,
                value=(f'=IF($B{r}="","",IFERROR(ROUND($K{r}*(1+Stammdaten!$B$17)'
                       f'*(1+Stammdaten!$B$19)/$E{r},2),0))')).number_format = EUR
        ws.cell(row=r, column=13, value=f'=IF($B{r}="","",ROUND($L{r}*$E{r},2))').number_format = EUR

        for col in range(1, 14):
            cell = ws.cell(row=r, column=col)
            cell.border = BOX
            if cell.font.color is None or cell.font.name != FONT:
                cell.font = CALC_F
            if i % 2:
                if cell.fill.fgColor.rgb != f"00{YELLOW}":
                    cell.fill = BAND_FILL

    dv = DataValidation(type="list", formula1=f"={kat_nr}", allow_blank=True, showDropDown=False)
    dv.error = "Bitte eine Katalog-Nr. aus dem Blatt 'Leistungskatalog' wählen."
    dv.prompt = "Katalog-Nr. wählen"
    ws.add_data_validation(dv)
    dv.add(f"B{POS_FIRST}:B{POS_LAST}")

    # --- Summen ----------------------------------------------------------
    s = POS_LAST + 2   # 46
    summen = [
        ("Zwischensumme Positionen (netto)", f"=ROUND(SUM($M${POS_FIRST}:$M${POS_LAST}),2)", False),
        ("An- und Abfahrt (km x Kilometersatz)",
         "=ROUND($B$21*Stammdaten!$B$20*(1+Stammdaten!$B$19),2)", False),
        ("Zwischensumme (netto)", f"=ROUND($M${s}+$M${s + 1},2)", False),
        ("Rabatt", f"=-ROUND($M${s + 2}*Stammdaten!$B$18,2)", False),
        ("Nettobetrag", f"=ROUND($M${s + 2}+$M${s + 3},2)", True),
        ("zzgl. Umsatzsteuer", f"=ROUND($M${s + 4}*Stammdaten!$B$21,2)", False),
        ("Gesamtbetrag (brutto)", f"=ROUND($M${s + 4}+$M${s + 5},2)", True),
    ]
    for i, (label, formula, bold) in enumerate(summen):
        r = s + i
        c = ws.cell(row=r, column=11, value=label)
        c.font = LABEL_B if bold else LABEL
        c.alignment = Alignment(horizontal="right")
        ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=12)
        v = ws.cell(row=r, column=13, value=formula)
        v.number_format = EUR
        v.font = Font(name=FONT, size=10, bold=bold)
        v.border = BOX
        if bold:
            v.fill = TOTAL_FILL
            c.fill = TOTAL_FILL

    # --- Kennzahlen ------------------------------------------------------
    k = s + 9   # 55
    section(ws, k - 1, "Kennzahlen (nur intern - nicht Teil des Angebots)", 13)
    netto = f"$M${s + 4}"
    kenn = [
        ("Gesamtstunden", f"=ROUND(SUM($G${POS_FIRST}:$G${POS_LAST}),2)", NUM2),
        ("Materialkosten (netto)", f"=ROUND(SUM($J${POS_FIRST}:$J${POS_LAST}),2)", EUR),
        # Fahrtkosten hier zum reinen Kostensatz, nicht zum Verkaufswert aus M{s+1}.
        ("Selbstkosten gesamt",
         f"=ROUND(SUM($K${POS_FIRST}:$K${POS_LAST})+$B$21*Stammdaten!$B$20,2)", EUR),
        ("Deckungsbeitrag", f"=ROUND({netto}-$B${k + 2},2)", EUR),
        ("Marge auf Nettoumsatz", f"=IFERROR($B${k + 3}/{netto},0)", PCT),
        ("Angebotspreis je m² Wohnfläche", f"=IFERROR(ROUND({netto}/$B$11,2),0)", EUR),
        ("Erlös je Arbeitsstunde", f"=IFERROR(ROUND({netto}/$B${k},2),0)", EUR),
    ]
    for i, (label, formula, fmt) in enumerate(kenn):
        r = k + i
        ws.cell(row=r, column=1, value=label).font = LABEL
        c = ws.cell(row=r, column=2, value=formula)
        c.number_format = fmt
        c.font = CALC_F
        c.border = BOX

    ws.freeze_panes = f"A{POS_FIRST}"
    return ws, s


# =========================================================================
def build_angebot(wb, sum_row):
    ws = wb.create_sheet("Angebot")
    set_widths(ws, {"A": 6, "B": 54, "C": 11, "D": 9, "E": 14, "F": 15})
    ws.sheet_view.showGridLines = False

    for r in range(1, 8):
        ws.row_dimensions[r].height = 15

    ws["A1"] = "=Stammdaten!$B$4"
    ws["A1"].font = Font(name=FONT, size=15, bold=True, color="1F3864")
    ws.merge_cells("A1:D1")
    for i, ref in enumerate(["=Stammdaten!$B$5", "=Stammdaten!$B$6",
                             '="Tel. "&Stammdaten!$B$7', "=Stammdaten!$B$8"]):
        c = ws.cell(row=2 + i, column=1, value=ref)
        c.font = LINK_F
        ws.merge_cells(start_row=2 + i, start_column=1, end_row=2 + i, end_column=3)

    ws["E1"] = "ANGEBOT"
    ws["E1"].font = Font(name=FONT, size=15, bold=True, color="1F3864")
    ws.merge_cells("E1:F1")
    ws["E1"].alignment = Alignment(horizontal="right")

    meta = [
        ("Angebotsnr.", "=Kalkulation!$B$7", None),
        ("Datum", "=Kalkulation!$B$8", "DD.MM.YYYY"),
        ("Gültig bis", "=Kalkulation!$B$8+Stammdaten!$B$23", "DD.MM.YYYY"),
    ]
    for i, (label, ref, fmt) in enumerate(meta):
        c = ws.cell(row=3 + i, column=5, value=label)
        c.font = LABEL
        c.alignment = Alignment(horizontal="right")
        v = ws.cell(row=3 + i, column=6, value=ref)
        v.font = LINK_F
        v.alignment = Alignment(horizontal="right")
        if fmt:
            v.number_format = fmt

    for i, ref in enumerate(["=Kalkulation!$B$3", "=Kalkulation!$B$4", "=Kalkulation!$B$5"]):
        c = ws.cell(row=8 + i, column=1, value=ref)
        c.font = Font(name=FONT, size=11, bold=(i == 0), color=GREEN)
        ws.merge_cells(start_row=8 + i, start_column=1, end_row=8 + i, end_column=3)

    ws["A13"] = '="Bauvorhaben: "&Kalkulation!$B$6'
    ws["A13"].font = Font(name=FONT, size=10, bold=True, color=GREEN)
    ws.merge_cells("A13:F13")
    ws["A14"] = ("Sehr geehrte Damen und Herren, vielen Dank für Ihre Anfrage. "
                 "Für die genannten Leistungen bieten wir Ihnen an:")
    ws["A14"].font = LABEL
    ws.merge_cells("A14:F14")

    head_row = 16
    for i, h in enumerate(["Pos", "Leistung", "Menge", "Einheit", "EP (€)", "Gesamt (€)"], start=1):
        c = ws.cell(row=head_row, column=i, value=h)
        c.font = H2
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center")

    first = head_row + 1
    for i in range(POS_ROWS):
        r = first + i
        src = POS_FIRST + i
        guard = f'IF(Kalkulation!$B${src}="",""'
        ws.cell(row=r, column=1, value=f'={guard},Kalkulation!$A${src})')
        ws.cell(row=r, column=2, value=f'={guard},Kalkulation!$C${src})')
        ws.cell(row=r, column=3, value=f'={guard},Kalkulation!$E${src})').number_format = NUM2
        ws.cell(row=r, column=4, value=f'={guard},Kalkulation!$D${src})')
        ws.cell(row=r, column=5, value=f'={guard},Kalkulation!$L${src})').number_format = EUR
        ws.cell(row=r, column=6, value=f'={guard},Kalkulation!$M${src})').number_format = EUR
        for col in range(1, 7):
            cell = ws.cell(row=r, column=col)
            cell.border = BOX
            cell.font = LINK_F
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")

    t = first + POS_ROWS + 1
    totals = [
        ("Zwischensumme Positionen", f"=Kalkulation!$M${sum_row}", False),
        ("An- und Abfahrt", f"=Kalkulation!$M${sum_row + 1}", False),
        ("Rabatt", f"=Kalkulation!$M${sum_row + 3}", False),
        ("Nettobetrag", f"=Kalkulation!$M${sum_row + 4}", True),
        ('=" zzgl. "&TEXT(Stammdaten!$B$21,"0 %")&" Umsatzsteuer"',
         f"=Kalkulation!$M${sum_row + 5}", False),
        ("Gesamtbetrag (brutto)", f"=Kalkulation!$M${sum_row + 6}", True),
    ]
    for i, (label, formula, bold) in enumerate(totals):
        r = t + i
        c = ws.cell(row=r, column=4, value=label)
        c.font = Font(name=FONT, size=10, bold=bold)
        c.alignment = Alignment(horizontal="right")
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        v = ws.cell(row=r, column=6, value=formula)
        v.number_format = EUR
        v.font = Font(name=FONT, size=11 if bold else 10, bold=bold,
                      color=BLACK if bold else GREEN)
        v.border = BOX
        if bold:
            v.fill = TOTAL_FILL
            c.fill = TOTAL_FILL
        if i == len(totals) - 1:
            v.border = Border(left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color="1F3864"))

    f = t + len(totals) + 2
    footer = [
        '="Zahlungsbedingungen: "&TEXT(Stammdaten!$B$22,"0")&" Tage netto, bei Zahlung innerhalb '
        'von 8 Tagen "&TEXT(Stammdaten!$B$19,"0 %")&" Skonto."',
        '="Dieses Angebot ist "&TEXT(Stammdaten!$B$23,"0")&" Tage ab Angebotsdatum gültig."',
        "Ausführungstermin nach Absprache. Es gelten unsere allgemeinen Geschäftsbedingungen.",
        '="USt-IdNr.: "&Stammdaten!$B$9&"   |   IBAN: "&Stammdaten!$B$10',
    ]
    for i, text in enumerate(footer):
        c = ws.cell(row=f + i, column=1, value=text)
        c.font = NOTE_F
        ws.merge_cells(start_row=f + i, start_column=1, end_row=f + i, end_column=6)
    ws.cell(row=f, column=1).border = TOP_LINE

    ws.cell(row=f + len(footer) + 2, column=1,
            value="Mit freundlichen Grüßen").font = LABEL
    ws.cell(row=f + len(footer) + 4, column=1, value="=Stammdaten!$B$4").font = LINK_F

    ws.print_area = f"A1:F{f + len(footer) + 4}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.6
    return ws


# =========================================================================
def main():
    out = sys.argv[1] if len(sys.argv) > 1 else \
        "assets/angebotskalkulator/Angebotskalkulator_Handwerk.xlsx"

    wb = Workbook()
    wb.remove(wb.active)

    build_anleitung(wb)
    build_stammdaten(wb)
    build_katalog(wb)
    _, sum_row = build_kalkulation(wb)
    build_angebot(wb, sum_row)

    wb.save(out)
    print(f"geschrieben: {out}")


if __name__ == "__main__":
    main()
